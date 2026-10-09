from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

import quotation as q

TODAY = date(2021, 7, 16)


def parse(form):
    return q.parse_form(form, today=TODAY)


def errors_of(form):
    with pytest.raises(q.QuotationError) as e:
        parse(form)
    return e.value.errors


def test_parse_valid_form(valid_form):
    quote = parse(valid_form)
    assert quote.customer_name == '江美志'
    assert quote.valid_until == date(2021, 8, 15)
    assert quote.ship_date == date(2021, 7, 22)
    assert [(p.name, p.quantity, p.price) for p in quote.products] == [
        ('電腦主機', 100, 5500), ('加購記憶體', 100, 400)]
    totals = quote.totals
    assert (totals.subtotal, totals.tax, totals.total) == (590000, 29500, 619500)


def test_tax_included_shows_included_tax(valid_form):
    totals = parse(dict(valid_form, tax='y')).totals
    assert totals.total == totals.subtotal == 590000   # 不另外加稅
    assert totals.tax == 28095                          # 590000 × 5 / 105 = 28095.24
    assert totals.net == 590000 - 28095


def P(*rows):
    return [q.Product(n, Decimal(qty), Decimal(price)) for n, qty, price in rows]


@pytest.mark.parametrize('products, tax_included, expected', [
    # 未稅：稅額四捨五入到元
    (P(('A', '1', '1234')), False, ('1234', '62', '1296', '1234')),
    (P(('A', '1', '1010')), False, ('1010', '51', '1061', '1010')),     # 50.5 → 51
    (P(('A', '1', '30')), False, ('30', '2', '32', '30')),              # 1.5 → 2
    (P(('A', '1', '29')), False, ('29', '1', '30', '29')),              # 1.45 → 1
    # 每列金額四捨五入到分，小計是四捨五入後的加總
    (P(('A', '1.5', '12.25')), False, ('18.38', '1', '19.38', '18.38')),
    (P(('A', '3', '0.333')), False, ('1.00', '0', '1.00', '1.00')),
    (P(('A', '2', '100'), ('折扣', '1', '-50')), False, ('150', '8', '158', '150')),
    # 含稅：從總額回推內含稅額，未稅金額 + 稅額 = 總額
    (P(('A', '1', '1050')), True, ('1050', '50', '1050', '1000')),
    (P(('A', '1', '1000')), True, ('1000', '48', '1000', '952')),
    (P(('A', '1', '105.5')), True, ('105.5', '5', '105.5', '100.5')),
    (P(('A', '1', '640.5')), True, ('640.5', '31', '640.5', '609.5')),   # 30.5 → 31
    (P(), False, ('0', '0', '0', '0')),
])
def test_compute_totals(products, tax_included, expected):
    totals = q.compute_totals(products, tax_included)
    assert (totals.subtotal, totals.tax, totals.total, totals.net) == tuple(Decimal(v) for v in expected)
    assert totals.net + totals.tax == totals.total if tax_included else totals.subtotal + totals.tax == totals.total


def test_tax_percent_is_whole_number():
    # 含稅公式用 *5/105，稅率必須是整數百分比
    assert q.TAX_PERCENT == q.TAX_RATE * 100
    assert q.TAX_INCLUDED_FORMULA == '=ROUND((G34+G37)*5/105,0)'


def test_quote_number_format():
    assert q.format_quote_no(date(2026, 10, 9), 7) == 'SD202610-007'
    assert q.format_quote_no(date(2026, 1, 1), 1234) == 'SD202601-1234'
    assert q.quote_no_prefix(date(2026, 12, 31)) == 'SD202612-'


def test_snapshot_round_trip(valid_form):
    quote = parse(dict(valid_form, tax='y', taxid='24268597', note='一\n二', product='A,1.50,12.25\n折扣,1,-5'))
    quote.quote_no = 'SD202107-001'
    assert q.Quotation.from_dict(quote.to_dict()) == quote
    quote.ship_date = None
    assert q.Quotation.from_dict(quote.to_dict()) == quote


def test_to_form_round_trip(valid_form):
    quote = parse(dict(valid_form, tax='y', note='一\n二', product='A,1.50,12.250\n折扣,1,-5\nB,1E+2,0'))
    again = parse(quote.to_form())
    assert again.products == quote.products
    assert again.ship_date is None  # 複製時不帶出貨日
    quote.ship_date = None
    assert again == quote


def test_preview():
    result = q.preview({'product': 'A,1,1234\nbad', 'tax': 'n'})
    assert len(result['errors']) == 1
    assert result['lines'] == ['1234.00']
    assert result['totals'].total == Decimal('1296')


@pytest.mark.parametrize('raw, expected', [
    ('A,1,2', [('A', 1, 2)]),
    ('A，1，2', [('A', 1, 2)]),
    ('A|B 規格,1,100', [('A|B 規格', 1, 100)]),
    ('  A , 1 , 2  ', [('A', 1, 2)]),
    ('A,1.5,12.25', [('A', Decimal('1.5'), Decimal('12.25'))]),
    ('A,1,$100', [('A', 1, 100)]),
    ('A,1,0', [('A', 1, 0)]),
    ('A,2,100\n折扣,1,-50', [('A', 2, 100), ('折扣', 1, -50)]),   # 單價負數 = 折扣列
    ('運費,1,150', [('運費', 1, 150)]),
    ('A,1,-0', [('A', 1, 0)]),
    ('A,1,95000000.1234', [('A', 1, Decimal('95000000.1234'))]),   # 含稅後 99,750,000 < 1 億
    ('A,1,2\n\n   \nB,3,4\n', [('A', 1, 2), ('B', 3, 4)]),
    ('A,1,2\r\nB,3,4', [('A', 1, 2), ('B', 3, 4)]),
])
def test_parse_products(raw, expected):
    errors = []
    products = q.parse_products(raw, errors)
    assert errors == []
    assert [(p.name, p.quantity, p.price) for p in products] == expected


@pytest.mark.parametrize('raw, message', [
    ('', '請至少輸入一個品項'),
    ('\n \n', '請至少輸入一個品項'),
    ('only name', '第 1 行'),
    ('A,1', '第 1 行'),
    ('A,1,2\nB,x,2', '第 2 行的數量'),
    ('A,0,2', '數量'),
    ('A,-1,2', '數量'),
    ('A,1,abc', '價格'),
    ('A,1,-5', '合計不可小於 0'),
    ('A,1e400,1', '數量'),
    ('A,1,1e999999', '價格'),
    ('A,1,1000000000000', '價格'),
    ('A,1,100000000', '價格'),                      # 單價要小於 1 億（再大報價單會印成 ###）
    ('A,100000000,1', '數量'),
    ('A,1000,100000', '金額太大'),                  # 每列金額 1 億
    ('A,1,99999999', '金額太大'),                   # 加稅後總計超過 1 億
    ('A,2,60000000\n折扣,1,-50000000', '金額太大'),  # 合計不大，但有一列超過 1 億
    ('A,1,0.00001', '價格'),
    ('A,NaN,5', '數量'),
    ('A,1,Infinity', '價格'),
    ('筆電,2,1,000', '不要寫成 1,000'),     # 千分位逗號不能默默變成單價 1
    ('主機, 含螢幕,2,100', '第 1 行'),
    (',1,2', '第 1 行'),
    ('中' * 43 + ',1,1', '商品名稱太長'),
])
def test_parse_products_errors(raw, message):
    errors = []
    q.parse_products(raw, errors)
    assert len(errors) == 1
    assert message in errors[0]


def test_too_many_products(valid_form):
    rows = '\n'.join(f'品項{i},1,1' for i in range(q.MAX_PRODUCTS + 1))
    assert any('最多' in e for e in errors_of(dict(valid_form, product=rows)))
    rows = '\n'.join(f'品項{i},1,1' for i in range(q.MAX_PRODUCTS))
    assert len(parse(dict(valid_form, product=rows)).products) == q.MAX_PRODUCTS


def test_missing_products_reported_with_other_errors(valid_form):
    errors = errors_of(dict(valid_form, cname='', product=''))
    assert '請填寫客戶姓名' in errors
    assert '請至少輸入一個品項' in errors


def test_all_errors_reported_together(valid_form):
    errors = errors_of(dict(valid_form, cname=' ', cphone='', taxid='1234', vday='abc',
                            seller='路人', delivery='飛機', cash='比特幣', tax='x', dday='明天'))
    assert len(errors) == 9


@pytest.mark.parametrize('taxid', ['1234567', '123456789', 'abcdefgh', '1234 5678'])
def test_invalid_tax_id(valid_form, taxid):
    assert '公司統編必須是 8 位數字' in errors_of(dict(valid_form, taxid=taxid))


def test_missing_optional_fields(valid_form):
    form = {k: v for k, v in valid_form.items() if k not in ('taxid', 'companyName', 'companyAddress', 'dday', 'note')}
    quote = parse(form)
    assert quote.ship_date is None and quote.tax_id == '' and quote.note == ''


def test_ship_date_accepts_slashes(valid_form):
    assert parse(dict(valid_form, dday='2021/07/22')).ship_date == date(2021, 7, 22)


def test_text_too_long(valid_form):
    assert any('不可超過' in e for e in errors_of(dict(valid_form, cname='x' * 201)))


@pytest.mark.parametrize('company, customer, expected', [
    ('測試股份有限公司', '江美志', '測試股份有限公司報價單0716'),
    ('', '江美志', '江美志報價單0716'),
    ('A/B\\C:D*E?F"G<H>I|J', '', 'ABCDEFGHIJ報價單0716'),
    ('../../etc', '', 'etc報價單0716'),
    ('...', '', '客戶報價單0716'),
])
def test_file_stem(valid_form, company, customer, expected):
    quote = parse(dict(valid_form, companyName=company, cname=customer or 'x'))
    quote.customer_name = customer
    assert quote.file_stem() == expected


def test_template_tax_rate_matches_constant():
    sheet = load_workbook(q.TEMPLATE_PATH)[q.SHEET_NAME]
    assert Decimal(str(sheet['G35'].value)) == q.TAX_RATE


def build_sheet(form, tmp_path):
    path = tmp_path / 'out.xlsx'
    q.build_workbook(parse(form)).save(path)
    wb = load_workbook(path)
    return wb, wb[q.SHEET_NAME]


def test_build_workbook_fills_template(valid_form, tmp_path):
    wb, sheet = build_sheet(dict(valid_form, taxid='24268597', note='貨到付款'), tmp_path)
    assert sheet['G3'].value == '2021/07/16'
    assert sheet['G8'].value == '2021/08/15'
    assert sheet['B9'].value == '姓名：江美志'
    assert sheet['B10'].value == '公司名稱：測試股份有限公司'
    assert sheet['B11'].value == '統一編號：24268597'
    assert sheet['B12'].value == '公司地址：新竹縣竹北市中和街1號'
    assert sheet['B13'].value == '公司電話：03-5520981#1688'
    assert sheet['C16'].value == '貨到付款'
    assert [sheet[c].value for c in ('B20', 'C20', 'D20', 'F20')] == ['陳聖尹', '2021/07/22', '郵局', '匯款']
    assert wb.calculation.fullCalcOnLoad


def test_build_workbook_products_are_numbers_and_formulas_kept(valid_form, tmp_path):
    _, sheet = build_sheet(valid_form, tmp_path)
    assert [sheet[f'{c}24'].value for c in 'BCEF'] == [100, '電腦主機', 5500, 'T']
    assert [sheet[f'{c}25'].value for c in 'BCEF'] == [100, '加購記憶體', 400, 'T']
    # 範本的範例列被清掉，沒用到的列是空的
    assert [sheet[f'{c}26'].value for c in 'BCEF'] == [None] * 4
    # 每列四捨五入到分、稅額四捨五入到元
    assert [sheet[f'G{r}'].value for r in (24, 33)] == ['=ROUND(B24*E24,2)', '=ROUND(B33*E33,2)']
    # 加總先四捨五入到分，浮點誤差才不會讓稅額少 1 元
    assert sheet['G34'].value == '=ROUND(SUM(G24:G33),2)'
    assert sheet['F36'].value == '稅額'
    assert sheet['G36'].value == '=ROUND(G35*ROUND(SUMIF(F24:F33,"T",G24:G33),2),0)'
    assert sheet['G38'].value == '=G34+G36+G37'
    # 換掉公式不會弄丟範本的金額格式（第一列有 $ 符號）
    assert '$' in sheet['G24'].number_format and '$' not in sheet['G25'].number_format
    assert '$' in sheet['G34'].number_format
    # 單價最多顯示 4 位小數（0.125 不會印成 0.13）
    assert all('#,##0.00##' in sheet[f'E{r}'].number_format for r in range(24, 34))
    assert '$' in sheet['E24'].number_format


def test_build_workbook_tax_included(valid_form, tmp_path):
    _, sheet = build_sheet(dict(valid_form, tax='y'), tmp_path)
    assert [sheet[f'F{r}'].value for r in range(24, 34)] == [None] * 10
    assert sheet['F36'].value == '內含稅額'
    assert sheet['G36'].value == '=ROUND((G34+G37)*5/105,0)'
    assert sheet['G38'].value == '=G34+G37'  # 稅已經含在小計裡，不能再加一次


def test_build_workbook_quote_number(valid_form, tmp_path):
    _, sheet = build_sheet(valid_form, tmp_path)
    assert sheet['F4'].value is None and sheet['G4'].value is None  # 還沒取號
    quote = parse(valid_form)
    quote.quote_no = 'SD202107-001'
    path = tmp_path / 'numbered.xlsx'
    q.build_workbook(quote).save(path)
    sheet = load_workbook(path)[q.SHEET_NAME]
    assert (sheet['F4'].value, sheet['G4'].value) == ('報價單號', 'SD202107-001')


@pytest.mark.parametrize('seller', q.SELLERS)
def test_contact_line(valid_form, tmp_path, seller):
    template_line = load_workbook(q.TEMPLATE_PATH)[q.SHEET_NAME]['B40'].value
    _, sheet = build_sheet(dict(valid_form, seller=seller), tmp_path)
    contact = q.SELLER_CONTACTS.get(seller)
    if contact:
        assert sheet['B40'].value == f'如您有任何疑問，請即聯絡：{seller}，{contact}'
    else:
        assert sheet['B40'].value == template_line  # 沒設定聯絡方式的銷售員保留範本原本的聯絡人


def test_contact_line_for_first_seller_unchanged(valid_form, tmp_path):
    _, sheet = build_sheet(dict(valid_form, seller='陳聖尹'), tmp_path)
    assert sheet['B40'].value == '如您有任何疑問，請即聯絡：陳聖尹，0931330086，teching_chen2000@yahoo.com.tw'


def test_build_workbook_decimal_numbers(valid_form, tmp_path):
    _, sheet = build_sheet(dict(valid_form, product='A,1.5,12.25'), tmp_path)
    assert (sheet['B24'].value, sheet['E24'].value) == (1.5, 12.25)


def test_build_workbook_without_company(valid_form, tmp_path):
    _, sheet = build_sheet(dict(valid_form, companyName='', companyAddress='', dday='', note=''), tmp_path)
    assert sheet['B10'].value is None
    assert sheet['B11'].value == '統一編號：無'
    assert sheet['B12'].value is None
    assert sheet['C16'].value == '無'
    assert sheet['C20'].value == '-'


def test_user_input_never_becomes_a_formula(valid_form, tmp_path):
    payload = '=HYPERLINK("http://evil")'
    _, sheet = build_sheet(dict(valid_form, note=payload, product=f'{payload},1,1'), tmp_path)
    for cell in (sheet['C16'], sheet['C24']):
        assert cell.value == payload
        assert cell.data_type == 's'


def test_build_workbook_keeps_stamp_image(valid_form, tmp_path):
    pytest.importorskip('PIL')  # 沒有 Pillow 時 openpyxl 會直接丟掉範本裡的圖片
    _, sheet = build_sheet(valid_form, tmp_path)
    assert len(sheet._images) == 1


def test_negative_zero_becomes_zero():
    errors = []
    product, = q.parse_products('A,1,-0.00', errors)
    assert not product.price.is_signed() and not product.amount.is_signed()


def test_control_characters_are_removed(valid_form):
    quote = parse(dict(valid_form, cname='江\x00美\uffff志\ufffe', note='第一行\x0b第二行\x0c第三行', product='A\x01B,1,2'))
    assert quote.customer_name == '江美志'
    assert quote.note == '第一行\n第二行\n第三行'
    assert quote.products[0].name == 'AB'


def test_note_must_fit_three_lines(valid_form):
    assert parse(dict(valid_form, note='一\n二\n三')).note == '一\n二\n三'
    assert any('備註太長' in e for e in errors_of(dict(valid_form, note='一\n二\n三\n四')))
    assert any('備註太長' in e for e in errors_of(dict(valid_form, note='字' * 120)))


def test_wrapped_lines():
    assert q.display_width('ab中文') == 6
    assert q.wrapped_lines('', 10) == 1
    assert q.wrapped_lines('a' * 10, 10) == 1
    assert q.wrapped_lines('a' * 11, 10) == 2
    assert q.wrapped_lines('中' * 6 + '\n\nx', 10) == 4


def test_build_workbook_note_and_long_names_wrap(valid_form, tmp_path):
    long_name = '工業用不鏽鋼六角螺絲 M3x10mm 304材質 (100入/包)'
    _, sheet = build_sheet(dict(valid_form, note='一\n二', product=f'{long_name},1,1\n短,1,1'), tmp_path)
    assert 'C16:G18' in {str(r) for r in sheet.merged_cells.ranges}
    assert sheet['C16'].value == '一\n二'
    assert sheet['C16'].alignment.wrap_text
    assert sheet['C24'].alignment.wrap_text
    assert sheet.row_dimensions[24].height == pytest.approx(20.1 * q.wrapped_lines(long_name, q.PRODUCT_NAME_LINE_WIDTH))
    assert sheet.row_dimensions[25].height == pytest.approx(20.1)
    assert not sheet['C25'].alignment.wrap_text
    assert sheet.sheet_properties.pageSetUpPr.fitToPage
    assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 1)


def test_full_width_tax_id_is_normalized(valid_form):
    assert parse(dict(valid_form, taxid='２４２６８５９７')).tax_id == '24268597'


def test_file_stem_uses_quote_number(valid_form):
    quote = parse(valid_form)
    quote.quote_no = 'SD202107-001'
    assert quote.file_stem() == '測試股份有限公司報價單SD202107-001'


def test_note_presets_fit_on_one_line():
    for preset in q.NOTE_PRESETS:
        assert q.wrapped_lines(preset, q.NOTE_LINE_WIDTH) == 1
